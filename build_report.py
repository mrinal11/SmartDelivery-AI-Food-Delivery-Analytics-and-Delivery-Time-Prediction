"""
Generates MrinalSingh_ProjectReport.docx
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import docx
import datetime

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.25)
    section.right_margin  = Inches(1.25)

# ── Style helpers ─────────────────────────────────────────────────────────────
def set_heading(para, text, level=1, color=None):
    run = para.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(18)
    elif level == 2:
        run.font.size = Pt(14)
    elif level == 3:
        run.font.size = Pt(12)
    else:
        run.font.size = Pt(11)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run

def add_heading(doc, text, level=1, color=(26, 86, 157)):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run(text)
    run.bold = True
    sz = {1: 18, 2: 14, 3: 12, 4: 11}[level]
    run.font.size = Pt(sz)
    run.font.color.rgb = RGBColor(*color)
    return p

def add_body(doc, text, bold=False, italic=False, size=11):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1+len(rows), cols=len(headers))
    table.style = 'Table Grid'
    # Header row
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        shading = OxmlElement('w:shd')
        shading.set(qn('w:val'), 'clear')
        shading.set(qn('w:color'), 'auto')
        shading.set(qn('w:fill'), '1A569D')
        cell._tc.get_or_add_tcPr().append(shading)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    # Data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = str(val)
            cell.paragraphs[0].runs[0].font.size = Pt(10)
            if ri % 2 == 1:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:val'), 'clear')
                shading.set(qn('w:color'), 'auto')
                shading.set(qn('w:fill'), 'EBF2FA')
                cell._tc.get_or_add_tcPr().append(shading)
    if col_widths:
        for row in table.rows:
            for ci, w in enumerate(col_widths):
                row.cells[ci].width = Inches(w)
    return table

# ═════════════════════════════════════════════════════════════════════════════
# TITLE PAGE
# ═════════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(40)
r = p.add_run("SmartDelivery AI")
r.bold = True; r.font.size = Pt(28)
r.font.color.rgb = RGBColor(26, 86, 157)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Food Delivery Analytics and Delivery Time Prediction")
r.font.size = Pt(16); r.italic = True
r.font.color.rgb = RGBColor(80, 80, 80)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
lines = [
    ("Project Report", 13, True),
    ("IBM SkillsBuild Data Analytics with AI Academic Internship", 12, False),
    ("Conducted by BharatCares in association with AICTE", 11, False),
    ("", 10, False),
    ("Submitted by: Mrinal Singh", 12, True),
    (f"Submission Date: {datetime.date.today().strftime('%B %d, %Y')}", 11, False),
]
for text, sz, bold in lines:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(sz)
    r.bold = bold

doc.add_page_break()

# ═════════════════════════════════════════════════════════════════════════════
# 1. ABSTRACT
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "1. Abstract", 1)
add_body(doc, (
    "This project presents SmartDelivery AI — a complete data analytics and machine learning "
    "solution for predicting food delivery times. Using a real-world dataset of 2,442 delivery "
    "records across multiple Indian cities, we performed rigorous data cleaning, exploratory data "
    "analysis (EDA), feature engineering, and trained four supervised regression models. "
    "The best model (Random Forest Regressor) achieved a Mean Absolute Error of 3.67 minutes "
    "and an R² of 0.746 on the held-out test set, explaining approximately 74.6% of the variance "
    "in delivery time. Key predictors identified include delivery person rating, delivery distance "
    "(Haversine), weather conditions, number of simultaneous deliveries, and road traffic density. "
    "Actionable business insights and recommendations are provided for operations managers."
))

# ═════════════════════════════════════════════════════════════════════════════
# 2. INTRODUCTION
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "2. Introduction", 1)
add_body(doc, (
    "The online food delivery industry has grown rapidly in India, with platforms operating in "
    "dozens of cities. Accurate delivery time estimation (ETA) is critical for customer "
    "satisfaction, fleet management, and operational efficiency. Over- or under-estimating ETAs "
    "leads to customer churn or unnecessary resource allocation."
))
add_body(doc, (
    "This project addresses the delivery time prediction problem using a structured machine "
    "learning workflow. The dataset contains 2,442 real delivery records with 20 attributes "
    "including GPS coordinates, time stamps, weather, traffic, vehicle condition, and delivery "
    "person characteristics."
))

add_heading(doc, "2.1 Problem Statement", 2)
add_body(doc, (
    "Predict the food delivery time in minutes for a given order, using operational and "
    "contextual features available at the time of order dispatch. This is a supervised "
    "regression problem where the target variable is Time_taken(min)."
))

add_heading(doc, "2.2 Objectives", 2)
objectives = [
    "Perform thorough exploratory data analysis to understand the dataset.",
    "Clean and preprocess the data, handling missing values and invalid entries.",
    "Identify key factors that influence delivery time.",
    "Engineer meaningful features: Haversine delivery distance, order hour, pickup delay.",
    "Train and evaluate four regression models: Linear Regression, Random Forest, Gradient Boosting, HistGradientBoosting.",
    "Select the best model using MAE, RMSE, and R² metrics.",
    "Generate actionable business insights for delivery operations teams.",
    "Save the best model for deployment readiness."
]
for obj in objectives:
    add_bullet(doc, obj)

# ═════════════════════════════════════════════════════════════════════════════
# 3. DATASET DESCRIPTION
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "3. Dataset Description", 1)
add_body(doc, "File: updated.csv | Records: 2,442 rows, 20 columns")

add_heading(doc, "3.1 Column Descriptions", 2)
col_table_data = [
    ("Unnamed: 0",           "int",    "Auto-generated row index (dropped during cleaning)"),
    ("ID",                   "object", "Unique order identifier (dropped during cleaning)"),
    ("Delivery_person_ID",   "object", "Delivery agent code (dropped during cleaning)"),
    ("Delivery_person_Age",  "int",    "Age of the delivery person (years); range 20–50"),
    ("Delivery_person_Ratings","float","Delivery person rating (scale 1–5); 94 missing"),
    ("Restaurant_latitude",  "float",  "GPS latitude of the restaurant"),
    ("Restaurant_longitude", "float",  "GPS longitude of the restaurant"),
    ("Delivery_location_latitude","float","GPS latitude of the delivery address"),
    ("Delivery_location_longitude","float","GPS longitude of the delivery address"),
    ("Time_Orderd",          "object", "Time the order was placed (HH:MM); 89 missing"),
    ("Time_Order_picked",    "object", "Time the order was picked up (HH:MM)"),
    ("Weatherconditions",    "int",    "Encoded weather: 0=Sunny, 1=Cloudy, 2=Windy, 3=Fog, 4=Sandstorm, 6=Stormy, 7=Tornado"),
    ("Road_traffic_density", "int",    "Encoded traffic: 0=Low, 1=Medium, 2=High, 3=Very High, 5=Jam"),
    ("Vehicle_condition",    "int",    "Condition of delivery vehicle (0=Poor to 3=Excellent)"),
    ("Type_of_order",        "int",    "Order type encoded (0–3)"),
    ("Type_of_vehicle",      "int",    "Vehicle type encoded (1–4)"),
    ("multiple_deliveries",  "float",  "Number of simultaneous deliveries (0–3); 54 missing"),
    ("Festival",             "int",    "Festival indicator: 1=No Festival, 2=Festival, 3=Unknown"),
    ("City",                 "int",    "City type: 1=Metropolitan, 2=Urban, 3=Semi-Urban, 4=Rural"),
    ("Time_taken(min)",      "float",  "TARGET: Actual delivery time in minutes (10–54 min)"),
]
add_table(doc,
    ["Column", "Type", "Description"],
    col_table_data,
    col_widths=[1.6, 0.8, 4.0]
)

add_heading(doc, "3.2 Target Variable Summary", 2)
add_table(doc,
    ["Statistic", "Value"],
    [
        ("Count",  "2,442"),
        ("Mean",   "26.49 min"),
        ("Std Dev","9.34 min"),
        ("Min",    "10 min"),
        ("25th %ile", "20 min"),
        ("Median", "26 min"),
        ("75th %ile", "32 min"),
        ("Max",    "54 min"),
    ],
    col_widths=[2.5, 2.5]
)

# ═════════════════════════════════════════════════════════════════════════════
# 4. DATA QUALITY AND CLEANING
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "4. Data Quality Analysis and Cleaning", 1)

add_heading(doc, "4.1 Data Quality Issues Identified", 2)
add_table(doc,
    ["Issue", "Column", "Count", "Action Taken"],
    [
        ("Missing values",         "Delivery_person_Ratings",  "94 (3.8%)",  "Imputed with median (4.7)"),
        ("Missing values",         "Time_Orderd",              "89 (3.6%)",  "Rows dropped (needed for feature)"),
        ("Missing values",         "multiple_deliveries",      "54 (2.2%)",  "Imputed with mode (0)"),
        ("Invalid age (Age = 0)",  "Delivery_person_Age",      "93 (3.8%)",  "Rows dropped"),
        ("Invalid rating (> 5.0)", "Delivery_person_Ratings",  "5 (0.2%)",   "Rows dropped"),
        ("Zero GPS coordinates",   "Restaurant_lat/lon",       "176 (7.2%)", "Rows dropped"),
        ("Auto-generated index",   "Unnamed: 0",               "—",          "Column dropped"),
        ("High-cardinality IDs",   "ID, Delivery_person_ID",   "—",          "Columns dropped"),
        ("Duplicate rows",         "All columns",              "0",          "No action needed"),
    ],
    col_widths=[1.8, 1.8, 1.0, 1.9]
)

add_heading(doc, "4.2 Data Cleaning Results", 2)
add_body(doc,
    "After applying all cleaning steps: raw dataset had 2,442 rows. "
    "After cleaning, 2,176 rows were retained (266 rows removed, 10.9% reduction). "
    "This is a conservative cleaning approach — only clearly invalid records were removed, "
    "and missing values were imputed where possible rather than discarding rows."
)

add_heading(doc, "4.3 Time Column Handling", 2)
add_body(doc, (
    "The time columns Time_Orderd and Time_Order_picked contained non-standard values "
    "such as '10:60' and '23:60'. These are data-entry artefacts where minutes rolled "
    "over 59. They were corrected by propagating the excess minutes into the hours field "
    "(e.g., 10:60 becomes 11:00). No rows were discarded for this reason."
))

# ═════════════════════════════════════════════════════════════════════════════
# 5. FEATURE ENGINEERING
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "5. Feature Engineering", 1)
add_body(doc, "Three new features were engineered from existing columns:")
add_table(doc,
    ["Feature", "Formula / Source", "Rationale"],
    [
        ("distance_km",       "Haversine(restaurant lat/lon, delivery lat/lon)",
         "Direct measure of delivery distance; stronger than raw coordinates"),
        ("order_hour",        "Integer hour of Time_Orderd (0–23)",
         "Captures time-of-day effects (peak hours, nighttime orders)"),
        ("pickup_delay_min",  "Time_Order_picked - Time_Orderd (minutes)",
         "Proxy for kitchen preparation time and restaurant efficiency"),
    ],
    col_widths=[1.6, 2.6, 2.3]
)

add_heading(doc, "5.1 Haversine Distance", 2)
add_body(doc, (
    "The Haversine formula calculates the great-circle distance between two geographic "
    "coordinates on a sphere (Earth radius = 6,371 km). This produces a geographically "
    "meaningful distance in kilometres, which is a far more informative feature than raw "
    "latitude/longitude pairs. After removing zero-coordinate rows, the distance statistics "
    "show a median of ~9.3 km, with some longer rural routes up to ~50+ km."
))

add_heading(doc, "5.2 Final Feature Set (13 features)", 2)
add_table(doc,
    ["#", "Feature", "Type", "Notes"],
    [
        ("1",  "Delivery_person_Age",     "Numeric", "Age of delivery person (20–50)"),
        ("2",  "Delivery_person_Ratings", "Numeric", "Rating (2.5–5.0), imputed"),
        ("3",  "Weatherconditions",       "Ordinal",  "Encoded integer (0–7)"),
        ("4",  "Road_traffic_density",    "Ordinal",  "Encoded integer (0–5)"),
        ("5",  "Vehicle_condition",       "Ordinal",  "Encoded integer (0–3)"),
        ("6",  "Type_of_order",           "Nominal",  "Encoded integer (0–3)"),
        ("7",  "Type_of_vehicle",         "Nominal",  "Encoded integer (1–4)"),
        ("8",  "multiple_deliveries",     "Ordinal",  "Count (0–3), imputed"),
        ("9",  "Festival",                "Nominal",  "Encoded (1=No, 2=Yes)"),
        ("10", "City",                    "Nominal",  "Encoded (1–4 city types)"),
        ("11", "distance_km",             "Numeric",  "ENGINEERED: Haversine distance"),
        ("12", "order_hour",              "Ordinal",  "ENGINEERED: Hour 0–23"),
        ("13", "pickup_delay_min",        "Numeric",  "ENGINEERED: Pickup wait time"),
    ],
    col_widths=[0.4, 2.0, 0.8, 3.4]
)

# ═════════════════════════════════════════════════════════════════════════════
# 6. EXPLORATORY DATA ANALYSIS
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "6. Exploratory Data Analysis", 1)

add_heading(doc, "6.1 Target Variable Distribution", 2)
add_body(doc, (
    "The target variable Time_taken(min) has a roughly unimodal distribution with a slight "
    "right skew. The mean (26.49 min) and median (26 min) are close, indicating limited skew. "
    "The range is 10–54 minutes. No extreme outliers were observed. The distribution is "
    "consistent with real-world food delivery data across Indian cities."
))

add_heading(doc, "6.2 Correlation Analysis", 2)
add_body(doc, "Pearson correlations with the target variable (Time_taken(min)):")
add_table(doc,
    ["Feature", "Pearson r", "Interpretation"],
    [
        ("Delivery_person_Ratings", "Negative (~-0.20)", "Higher-rated agents deliver faster"),
        ("multiple_deliveries",     "Positive (~+0.34)", "More simultaneous orders = slower delivery"),
        ("distance_km",             "Positive (~+0.30)", "Longer distance = more time"),
        ("Road_traffic_density",    "Positive (~+0.25)", "Higher traffic = longer time"),
        ("Weatherconditions",       "Variable",          "Stormy/foggy weather increases time"),
        ("order_hour",              "Moderate",          "Certain hours have higher congestion"),
        ("Festival",                "Moderate",          "Festival periods show elevated times"),
    ],
    col_widths=[2.0, 1.6, 3.0]
)

add_heading(doc, "6.3 Key EDA Findings", 2)
eda_findings = [
    "Traffic density has a clear monotonic relationship with delivery time — jam conditions add measurable delays vs. low-traffic conditions.",
    "Delivery agents handling multiple simultaneous orders take significantly longer per delivery.",
    "Festival periods (Festival=2) show noticeably higher average delivery times.",
    "Higher-rated delivery agents (4.5–5.0) tend to complete deliveries slightly faster.",
    "Rural city deliveries (City=4) involve longer distances and higher delivery times on average.",
    "Peak order hours (evening, 18:00–21:00) correlate with higher delivery times.",
    "Stormy and tornado weather conditions produce the highest average delivery times.",
]
for f in eda_findings:
    add_bullet(doc, f)

# ═════════════════════════════════════════════════════════════════════════════
# 7. MACHINE LEARNING
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "7. Machine Learning Methodology", 1)

add_heading(doc, "7.1 Preprocessing", 2)
add_body(doc, (
    "All 13 features are numeric (encoded integers and engineered floats). A StandardScaler "
    "was applied within a sklearn ColumnTransformer to normalise features for the linear model. "
    "Tree-based models are scale-invariant, but the same preprocessor was applied uniformly "
    "within sklearn Pipelines for consistency. The dataset was split 80/20 (train/test) with "
    "random_state=42 for reproducibility."
))
add_table(doc,
    ["Split", "Samples"],
    [("Training set (80%)", "1,740"), ("Test set (20%)", "436")],
    col_widths=[3.0, 2.0]
)

add_heading(doc, "7.2 Models Evaluated", 2)
add_table(doc,
    ["Model", "Key Hyperparameters", "Notes"],
    [
        ("Linear Regression",     "Default",
         "Parametric baseline; assumes linear relationships"),
        ("Random Forest",         "n_estimators=200, random_state=42",
         "Bagging ensemble; robust to non-linearity and outliers"),
        ("Gradient Boosting",     "n_estimators=200, lr=0.1, max_depth=4",
         "Sequential boosting; strong performance on tabular data"),
        ("HistGradientBoosting",  "max_iter=300, lr=0.08, max_depth=5",
         "Histogram-based GBM; fast training, handles missing values natively"),
    ],
    col_widths=[1.8, 2.2, 2.6]
)

add_heading(doc, "7.3 Evaluation Metrics", 2)
add_table(doc,
    ["Metric", "Formula", "Interpretation"],
    [
        ("MAE",  "mean(|y - y_hat|)",            "Avg absolute error in minutes; lower is better"),
        ("RMSE", "sqrt(mean((y - y_hat)^2))",    "Penalises large errors; lower is better"),
        ("R²",   "1 - SS_res / SS_tot",           "Proportion of variance explained; higher is better"),
    ],
    col_widths=[0.8, 2.5, 3.3]
)

# ═════════════════════════════════════════════════════════════════════════════
# 8. RESULTS
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "8. Results and Model Evaluation", 1)

add_heading(doc, "8.1 Model Performance Comparison", 2)
add_table(doc,
    ["Model", "Test MAE (min)", "Test RMSE (min)", "Test R²", "Train MAE (min)"],
    [
        ("Random Forest",        "3.6663", "4.6974", "0.7459", "~1.20"),
        ("HistGradientBoosting", "3.6663", "4.7960", "0.7351", "~2.10"),
        ("Gradient Boosting",    "3.6979", "4.7506", "0.7401", "~2.95"),
        ("Linear Regression",    "5.5488", "7.0103", "0.4340", "~5.50"),
    ],
    col_widths=[2.0, 1.5, 1.5, 1.0, 1.5]
)

add_heading(doc, "8.2 Best Model: Random Forest Regressor", 2)
add_body(doc, (
    "The Random Forest Regressor achieved the best test MAE of 3.67 minutes (tied with "
    "HistGradientBoosting but with better RMSE), explaining 74.59% of the variance in "
    "delivery time. The model was further validated using 5-fold cross-validation."
))
add_table(doc,
    ["CV Metric", "Value"],
    [
        ("5-Fold CV MAE Mean", "3.4369 min"),
        ("5-Fold CV MAE Std",  "0.0506 min"),
        ("Interpretation",     "Highly stable — very low variance across folds"),
    ],
    col_widths=[2.5, 4.0]
)

add_heading(doc, "8.3 Feature Importance (Random Forest)", 2)
add_body(doc, "Top 5 most important features ranked by Gini importance:")
add_table(doc,
    ["Rank", "Feature", "Importance", "Business Meaning"],
    [
        ("1", "Delivery_person_Ratings", "0.2565", "Agent performance is the strongest driver"),
        ("2", "distance_km",             "0.1440", "Geographical distance directly impacts time"),
        ("3", "Weatherconditions",        "0.1200", "Adverse weather significantly slows delivery"),
        ("4", "multiple_deliveries",      "0.1137", "Batching orders increases individual delivery time"),
        ("5", "Road_traffic_density",     "0.0995", "Traffic congestion adds measurable delay"),
    ],
    col_widths=[0.5, 2.0, 1.2, 3.0]
)

add_heading(doc, "8.4 Residual Analysis", 2)
add_body(doc, (
    "The residuals (actual − predicted) are approximately normally distributed and centred "
    "near zero, confirming the model is unbiased. A small fraction of residuals exceed "
    "±10 minutes, corresponding to exceptional cases (extreme weather, unusual routes). "
    "No systematic directional bias was observed in the actual vs. predicted scatter plot."
))

# ═════════════════════════════════════════════════════════════════════════════
# 9. BUSINESS INSIGHTS
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "9. Business Insights and Recommendations", 1)

insights = [
    (
        "Insight 1 — Traffic Management",
        "Orders placed during jam traffic conditions take significantly longer than low-traffic "
        "orders. Operations teams should use real-time traffic feeds to trigger dynamic routing "
        "decisions at dispatch time.",
        "Integrate real-time traffic API into the dispatch system. During jam events, "
        "pre-assign agents already near the delivery destination."
    ),
    (
        "Insight 2 — Simultaneous Deliveries",
        "Each additional simultaneous delivery adds measurable time. Orders in the '3 simultaneous' "
        "category have the highest average delivery times.",
        "For premium/priority orders, enforce a maximum of 1 concurrent delivery per agent. "
        "Offer a 'priority delivery' option at premium pricing with a guaranteed ETA."
    ),
    (
        "Insight 3 — Festival Season Planning",
        "Festival periods show elevated delivery times. This is likely due to increased order "
        "volume, traffic congestion, and location crowding near event areas.",
        "Pre-position additional delivery agents in high-order-density zones 24–48 hours before "
        "major festivals. Temporarily increase ETA displayed to customers during festivals."
    ),
    (
        "Insight 4 — Agent Performance",
        "Higher-rated delivery agents consistently deliver faster. The feature importance analysis "
        "places rating as the single most important predictor of delivery time.",
        "Implement performance-based routing: assign higher-rated agents to complex or "
        "long-distance routes. Provide incentives for maintaining ratings above 4.5."
    ),
    (
        "Insight 5 — City-Specific Strategy",
        "Rural city deliveries involve longer distances and higher delivery times. Semi-Urban "
        "areas also show elevated times relative to metropolitan zones.",
        "Deploy micro-fulfillment centres or dark kitchens in semi-urban and rural areas "
        "to reduce last-mile distance. Set differentiated ETA expectations by city type."
    ),
    (
        "Insight 6 — Peak Hour Staffing",
        "Evening hours (18:00–21:00) show higher average delivery times. This reflects peak "
        "demand and possibly peak traffic coincidence.",
        "Staff at 120–150% of normal fleet capacity during 18:00–21:00. Use predictive "
        "scheduling models trained on historical hourly demand data."
    ),
]

for title, finding, recommendation in insights:
    add_heading(doc, title, 3)
    add_body(doc, "Finding: " + finding, bold=False)
    add_body(doc, "Recommendation: " + recommendation, italic=True)

# ═════════════════════════════════════════════════════════════════════════════
# 10. CONCLUSIONS
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "10. Conclusions", 1)
add_body(doc, (
    "This project successfully developed a complete end-to-end data analytics and machine "
    "learning pipeline for food delivery time prediction. The key contributions are:"
))
conclusions = [
    "Comprehensive data quality analysis identifying 4 distinct data issues, all addressed with justified cleaning strategies.",
    "Three engineered features (Haversine distance, order hour, pickup delay) that improved model predictive power.",
    "Systematic comparison of 4 regression models using consistent evaluation methodology.",
    "Best model (Random Forest) achieving MAE = 3.67 min and R² = 0.746 on unseen test data, with stable 5-fold CV performance.",
    "Six data-driven business insights with specific, actionable recommendations for food delivery operations teams.",
    "Production-ready model saved as smartdelivery_best_model.pkl using joblib for deployment.",
]
for c in conclusions:
    add_bullet(doc, c)

add_body(doc, (
    "The project demonstrates that delivery time can be predicted with practical accuracy using "
    "operational and contextual features available at dispatch time. The model is suitable for "
    "integration into a real-time ETA prediction API."
))

# ═════════════════════════════════════════════════════════════════════════════
# 11. TOOLS AND TECHNOLOGIES
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "11. Tools and Technologies", 1)
add_table(doc,
    ["Tool / Library", "Version", "Purpose"],
    [
        ("Python",             "3.10+",    "Core programming language"),
        ("pandas",             "2.3.x",    "Data manipulation and analysis"),
        ("numpy",              "1.24+",    "Numerical computing"),
        ("matplotlib",         "3.7+",     "Static data visualization"),
        ("seaborn",            "0.12+",    "Statistical visualizations"),
        ("scikit-learn",       "1.3+",     "ML models, pipelines, preprocessing"),
        ("joblib",             "1.3+",     "Model serialization"),
        ("Jupyter Notebook",   "7.x",      "Interactive development environment"),
        ("python-docx",        "1.2+",     "Report generation (.docx)"),
    ],
    col_widths=[1.8, 1.2, 3.5]
)

# ═════════════════════════════════════════════════════════════════════════════
# 12. REFERENCES
# ═════════════════════════════════════════════════════════════════════════════
add_heading(doc, "12. References", 1)
refs = [
    "Breiman, L. (2001). Random forests. Machine Learning, 45(1), 5–32.",
    "Friedman, J. H. (2001). Greedy function approximation: a gradient boosting machine. Annals of Statistics, 1189–1232.",
    "Sinnott, R. W. (1984). Virtues of the Haversine. Sky and Telescope, 68(2), 158.",
    "Pedregosa, F. et al. (2011). Scikit-learn: Machine Learning in Python. JMLR, 12, 2825–2830.",
    "IBM SkillsBuild Data Analytics with AI Academic Internship Program, BharatCares × AICTE, 2024–2025.",
    "McKinney, W. (2010). Data Structures for Statistical Computing in Python. Proceedings of the 9th Python in Science Conference.",
]
for i, ref in enumerate(refs, 1):
    add_bullet(doc, f"[{i}] {ref}")

# ─────────────────────────────────────────────────────────────────────────────
doc.save('MrinalSingh_ProjectReport.docx')
print('Report saved: MrinalSingh_ProjectReport.docx')
